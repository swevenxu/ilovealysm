"""DOCX extractor using python-docx for Word document extraction."""

from __future__ import annotations

from pathlib import Path

import docx

from ..models import (
    ExtractionMethod,
    ExtractedPage,
    ExtractionResult,
    FileType,
    PageReport,
    VerificationReport,
    VerificationStatus,
)

# Approximate characters per "page" for DOCX (since DOCX doesn't have real pages)
CHARS_PER_PAGE = 3000


class DOCXExtractor:
    """Extract and verify DOCX (Word) files."""

    @staticmethod
    def verify(file_path: str | Path) -> VerificationReport:
        """Run verification checks on a DOCX file.

        Checks:
        1. Can the file be opened/parsed?
        2. Can paragraphs/headings be extracted?
        3. Estimated page count for consistency reporting.
        """
        filename = Path(file_path).name

        try:
            document = docx.Document(str(file_path))
        except Exception as e:
            return VerificationReport(
                filename=filename,
                file_type=FileType.DOCX,
                status=VerificationStatus.CORRUPTED,
                message=f"Cannot open DOCX: {e}",
                error=str(e),
            )

        try:
            paragraphs = document.paragraphs
            total_text = ""
            has_headings = False
            has_paragraphs = False

            for para in paragraphs:
                text = para.text.strip()
                if text:
                    total_text += text + "\n"
                    has_paragraphs = True
                    if para.style and para.style.name and "Heading" in para.style.name:
                        has_headings = True

            if not has_paragraphs:
                return VerificationReport(
                    filename=filename,
                    file_type=FileType.DOCX,
                    status=VerificationStatus.CORRUPTED,
                    message="DOCX file contains no extractable text or paragraphs.",
                )

            # Estimate page count based on character length
            estimated_pages = max(1, len(total_text) // CHARS_PER_PAGE)

            # Create page reports for estimated pages
            page_reports = [
                PageReport(
                    page_number=i + 1,
                    has_text_layer=True,
                    text_length=CHARS_PER_PAGE if i < estimated_pages - 1 else len(total_text) % CHARS_PER_PAGE,
                )
                for i in range(estimated_pages)
            ]

            message = (
                f"DOCX readable. ~{estimated_pages} estimated pages, "
                f"{len(paragraphs)} paragraphs"
            )
            if has_headings:
                message += ", headings detected"
            message += "."

            return VerificationReport(
                filename=filename,
                file_type=FileType.DOCX,
                status=VerificationStatus.READABLE,
                page_count=estimated_pages,
                pages=page_reports,
                message=message,
            )

        except Exception as e:
            return VerificationReport(
                filename=filename,
                file_type=FileType.DOCX,
                status=VerificationStatus.CORRUPTED,
                message=f"Error reading DOCX structure: {e}",
                error=str(e),
            )

    @staticmethod
    def extract(file_path: str | Path) -> ExtractionResult:
        """Extract text from a DOCX file.

        Extracts all paragraphs with their styles preserved as markdown-like formatting.
        Chunks output into ~page-sized segments for consistency with the page model.
        """
        filename = Path(file_path).name

        try:
            document = docx.Document(str(file_path))
        except Exception as e:
            return ExtractionResult(
                filename=filename,
                file_type=FileType.DOCX,
                error=f"Cannot open DOCX: {e}",
            )

        # Build full text with markdown-style formatting
        lines: list[str] = []

        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                lines.append("")
                continue

            style_name = para.style.name if para.style else ""

            # Apply heading formatting
            if "Heading 1" in style_name:
                lines.append(f"# {text}")
            elif "Heading 2" in style_name:
                lines.append(f"## {text}")
            elif "Heading 3" in style_name:
                lines.append(f"### {text}")
            elif "Heading 4" in style_name:
                lines.append(f"#### {text}")
            elif "List" in style_name:
                lines.append(f"- {text}")
            else:
                lines.append(text)

        # Also extract tables
        for table in document.tables:
            lines.append("")
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                lines.append(" | ".join(cells))
            lines.append("")

        full_text = "\n".join(lines)

        # Chunk into page-sized segments
        pages: list[ExtractedPage] = []
        text_lines = full_text.split("\n")
        current_chunk: list[str] = []
        current_chars = 0
        page_num = 1

        for line in text_lines:
            current_chunk.append(line)
            current_chars += len(line) + 1  # +1 for newline

            if current_chars >= CHARS_PER_PAGE:
                page_text = "\n".join(current_chunk)
                pages.append(
                    ExtractedPage(
                        page_number=page_num,
                        text=page_text,
                        extraction_method=ExtractionMethod.DIRECT,
                        char_count=len(page_text),
                    )
                )
                page_num += 1
                current_chunk = []
                current_chars = 0

        # Don't forget the last chunk
        if current_chunk:
            page_text = "\n".join(current_chunk)
            pages.append(
                ExtractedPage(
                    page_number=page_num,
                    text=page_text,
                    extraction_method=ExtractionMethod.DIRECT,
                    char_count=len(page_text),
                )
            )

        total_chars = sum(p.char_count for p in pages)

        return ExtractionResult(
            filename=filename,
            file_type=FileType.DOCX,
            page_count=len(pages),
            pages=pages,
            total_chars=total_chars,
        )
